"""
[INPUT] ICD-10 Beijing clinical edition XLSX and validated CDN samples.
[OUTPUT] Fixed ICD candidate catalog, deterministic code mappings and coverage reports.
[POS] Sole standard-name and standard-code authority for Phase 1.
[UPDATE] Keep SPEC_PHASE1 mapping and OOV policies synchronized with this module.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from mednorm.data import CDNSample
from mednorm.text import normalize_label, normalize_text


@dataclass(frozen=True, slots=True)
class ICDCatalog:
    names: tuple[str, ...]
    name_to_codes: Mapping[str, tuple[str, ...]]
    source_row_count: int

    def codes_for(self, name: str) -> tuple[str, ...]:
        return self.name_to_codes.get(normalize_label(name), ())

    def primary_code_for(self, name: str) -> str | None:
        codes = self.codes_for(name)
        return codes[0] if codes else None


@dataclass(frozen=True, slots=True)
class CatalogCoverage:
    sample_count: int
    total_label_occurrences: int
    covered_label_occurrences: int
    fully_covered_samples: int
    unique_gold_names: int
    covered_unique_gold_names: int
    oov_names: tuple[str, ...]

    @property
    def label_coverage(self) -> float:
        if self.total_label_occurrences == 0:
            return 0.0
        return self.covered_label_occurrences / self.total_label_occurrences

    @property
    def fully_covered_sample_rate(self) -> float:
        if self.sample_count == 0:
            return 0.0
        return self.fully_covered_samples / self.sample_count


def load_icd_catalog(path: Path) -> ICDCatalog:
    """Load the headerless official ICD workbook and preserve source ordering."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise ValueError(f"cannot read ICD workbook {path}: {exc}") from exc

    sheet = workbook.active
    names: list[str] = []
    code_lists: dict[str, list[str]] = {}
    source_row_count = 0
    try:
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if not row or all(value is None for value in row):
                continue
            if len(row) < 2 or row[0] is None or row[1] is None:
                raise ValueError(f"ICD row {row_index} must contain code and name")
            code = normalize_text(str(row[0]))
            name = normalize_label(str(row[1]))
            if not code or not name:
                raise ValueError(f"ICD row {row_index} has an empty code or name")
            source_row_count += 1
            if name not in code_lists:
                code_lists[name] = []
                names.append(name)
            if code not in code_lists[name]:
                code_lists[name].append(code)
    finally:
        workbook.close()

    if not names:
        raise ValueError(f"ICD workbook {path} contains no valid rows")
    name_to_codes = {name: tuple(codes) for name, codes in code_lists.items()}
    return ICDCatalog(tuple(names), name_to_codes, source_row_count)


def catalog_coverage(samples: Sequence[CDNSample], catalog: ICDCatalog) -> CatalogCoverage:
    """Measure coverage without dropping out-of-catalog gold labels."""
    candidate_names = set(catalog.names)
    total_labels = 0
    covered_labels = 0
    fully_covered_samples = 0
    unique_gold: set[str] = set()
    covered_unique: set[str] = set()
    oov_seen: set[str] = set()
    oov_names: list[str] = []

    for sample in samples:
        labels = tuple(dict.fromkeys(sample.labels))
        total_labels += len(labels)
        sample_is_covered = bool(labels)
        for label in labels:
            unique_gold.add(label)
            if label in candidate_names:
                covered_labels += 1
                covered_unique.add(label)
            else:
                sample_is_covered = False
                if label not in oov_seen:
                    oov_seen.add(label)
                    oov_names.append(label)
        if sample_is_covered:
            fully_covered_samples += 1

    return CatalogCoverage(
        sample_count=len(samples),
        total_label_occurrences=total_labels,
        covered_label_occurrences=covered_labels,
        fully_covered_samples=fully_covered_samples,
        unique_gold_names=len(unique_gold),
        covered_unique_gold_names=len(covered_unique),
        oov_names=tuple(oov_names),
    )

