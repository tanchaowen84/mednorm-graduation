from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import build_thesis as base


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "assets_v2"

base.SOURCE = ROOT / "论文正文_v2_简化图解版.md"
base.OUTPUT = ROOT / "张艺馨本科毕业论文_正式学术版.docx"
base.ASSET_DIR = ASSET_DIR

# The reference thesis uses a compact five-chapter structure and frequent
# explanatory figures. Keep the existing A4 thesis typography, but slightly
# reduce heading density and preserve a comfortable 1.5-line body rhythm.
base.TOKENS["h1_size_pt"] = 16
base.TOKENS["h2_size_pt"] = 13.5
base.TOKENS["h2_before_pt"] = 10
base.TOKENS["h2_after_pt"] = 6


def _font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont:
    regular = ROOT / "assets" / "fonts" / "NotoSansCJKsc-Regular.otf"
    bold_path = "/System/Library/Fonts/STHeiti Light.ttc"
    return ImageFont.truetype(bold_path if bold else str(regular), size=size)


def _text_center(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    font: ImageFont.FreeTypeFont,
    *,
    fill: str = "#1F2933",
    spacing: int = 10,
) -> None:
    left, top, right, bottom = box
    bounds = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
    width = bounds[2] - bounds[0]
    height = bounds[3] - bounds[1]
    draw.multiline_text(
        ((left + right - width) / 2, (top + bottom - height) / 2 - bounds[1]),
        text,
        font=font,
        fill=fill,
        spacing=spacing,
        align="center",
    )


def _rounded_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    *,
    fill: str,
    font: ImageFont.FreeTypeFont,
    outline: str = "#51606F",
    text_fill: str = "#1F2933",
    radius: int = 28,
    width: int = 4,
) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)
    _text_center(draw, box, text, font, fill=text_fill)


def _arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    *,
    fill: str = "#677585",
    width: int = 6,
) -> None:
    draw.line([start, end], fill=fill, width=width)
    sx, sy = start
    ex, ey = end
    dx, dy = ex - sx, ey - sy
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    base_x, base_y = ex - ux * 25, ey - uy * 25
    draw.polygon(
        [
            (ex, ey),
            (base_x + px * 12, base_y + py * 12),
            (base_x - px * 12, base_y - py * 12),
        ],
        fill=fill,
    )


def _new_canvas(width: int = 1800, height: int = 980) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (width, height), "white")
    return image, ImageDraw.Draw(image)


def _save(image: Image.Image, name: str) -> None:
    image.save(ASSET_DIR / name, dpi=(200, 200))


def generate_figures_v2() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    title_font = _font(48, bold=True)
    box_font = _font(38, bold=True)
    body_font = _font(31)
    small_font = _font(27)

    # Figure 1-1: one concrete normalization example.
    image, draw = _new_canvas(1800, 820)
    draw.text((70, 55), "原始诊断名称", font=title_font, fill="#1F2933")
    draw.text((1260, 55), "ICD 标准实体", font=title_font, fill="#1F2933")
    _rounded_box(draw, (80, 250, 540, 560), "右下叶肺癌", fill="#EAF2F8", font=_font(48, bold=True))
    _rounded_box(draw, (680, 270, 1110, 540), "疾病实体\n标准化", fill="#FFF3CD", font=box_font)
    _rounded_box(draw, (1240, 175, 1720, 390), "右肺下叶恶性肿瘤", fill="#DFF2E1", font=box_font)
    _rounded_box(draw, (1240, 470, 1720, 685), "癌", fill="#DFF2E1", font=_font(44, bold=True))
    _arrow(draw, (540, 405), (680, 405))
    _arrow(draw, (1110, 405), (1240, 280))
    _arrow(draw, (1110, 405), (1240, 575))
    draw.text((650, 700), "标准化结果由一个或多个标准实体组成", font=body_font, fill="#56616D")
    _save(image, "standardization_example.png")

    # Figure 1-2: normalization versus diagnosis.
    image, draw = _new_canvas(1800, 930)
    draw.text((90, 55), "疾病名称标准化", font=title_font, fill="#1F5F8B")
    _rounded_box(draw, (100, 170, 580, 410), "肺部恶性肿瘤", fill="#EAF2F8", font=box_font)
    _rounded_box(draw, (690, 170, 1110, 410), "统一名称", fill="#FFF3CD", font=box_font)
    _rounded_box(draw, (1220, 170, 1700, 410), "肺癌 / ICD 名称", fill="#DFF2E1", font=box_font)
    _arrow(draw, (580, 290), (690, 290))
    _arrow(draw, (1110, 290), (1220, 290))
    draw.line((100, 475, 1700, 475), fill="#CBD2D9", width=3)
    draw.text((90, 530), "症状辅助诊断", font=title_font, fill="#9A5A15")
    _rounded_box(draw, (100, 650, 650, 865), "一直咳嗽、夜间加重\n没有发烧", fill="#FDEBD0", font=body_font)
    _rounded_box(draw, (760, 650, 1120, 865), "临床推断", fill="#F8D7DA", font=box_font)
    _rounded_box(draw, (1230, 650, 1700, 865), "可能的疾病", fill="#FDE2E4", font=box_font)
    _arrow(draw, (650, 755), (760, 755), fill="#A8651D")
    _arrow(draw, (1120, 755), (1230, 755), fill="#A8651D")
    _save(image, "standardization_vs_diagnosis.png")

    # Figure 2-1: roles of the three datasets.
    image, draw = _new_canvas(1800, 940)
    sources = [
        ((80, 140, 520, 390), "CHIP-CDN\n原始诊断名称 + 标准答案", "#EAF2F8"),
        ((680, 140, 1120, 390), "ICD-10 v601\n37645 个候选名称", "#DFF2E1"),
        ((1280, 140, 1720, 390), "CPubMed-KG\n别名 + 医学关系", "#F3E8FF"),
    ]
    for box, text, fill in sources:
        _rounded_box(draw, box, text, fill=fill, font=body_font)
    targets = [
        ((150, 610, 650, 850), "模型训练与评价\n学习名称对应关系", "#E8F1FB"),
        ((650, 610, 1150, 850), "标准候选词表\n限定输出名称范围", "#E5F4E7"),
        ((1150, 610, 1650, 850), "外部知识增强\n补充别名与关系", "#F4EAFB"),
    ]
    for box, text, fill in targets:
        _rounded_box(draw, box, text, fill=fill, font=body_font)
    for start, end in [((300, 390), (390, 610)), ((900, 390), (900, 610)), ((1500, 390), (1410, 610))]:
        _arrow(draw, start, end)
    _save(image, "data_roles.png")

    # Figure 2-2: four common difficulty types.
    image, draw = _new_canvas(1800, 1120)
    cards = [
        ((80, 90, 850, 500), "简称与中英文混写", "卵巢 Ca\n→ 卵巢恶性肿瘤；癌", "#EAF2F8"),
        ((950, 90, 1720, 500), "部位与词序变化", "右下叶肺癌\n→ 右肺下叶恶性肿瘤；癌", "#DFF2E1"),
        ((80, 610, 850, 1020), "多个概念连写", "肺癌化疗贫血\n→ 肺恶性肿瘤；化学治疗；贫血；癌", "#FFF3CD"),
        ((950, 610, 1720, 1020), "否定或排除信息", "肺结核（-）\n→ 标准化结果不包含肺结核", "#FDE2E4"),
    ]
    for box, heading, example, fill in cards:
        draw.rounded_rectangle(box, radius=30, fill=fill, outline="#6B7785", width=4)
        left, top, right, _ = box
        draw.text((left + 40, top + 40), heading, font=box_font, fill="#1F2933")
        draw.line((left + 40, top + 115, right - 40, top + 115), fill="#AAB4BF", width=3)
        draw.multiline_text((left + 40, top + 165), example, font=body_font, fill="#34404C", spacing=18)
    _save(image, "difficulty_examples.png")

    # Figure 3-1: the simplified four-step method.
    image, draw = _new_canvas(1900, 900)
    steps = [
        ((60, 220, 430, 600), "1", "候选召回", "37645 → 400", "#EAF2F8"),
        ((540, 220, 910, 600), "2", "候选精排", "三类模型 + 两类证据", "#DFF2E1"),
        ((1020, 220, 1390, 600), "3", "数量预测", "分类模型 + 结构特征", "#FFF3CD"),
        ((1500, 220, 1870, 600), "4", "约束解码", "候选词表 + 分组阈值", "#F3E8FF"),
    ]
    for box, num, heading, detail, fill in steps:
        draw.rounded_rectangle(box, radius=30, fill=fill, outline="#596777", width=4)
        left, top, right, _ = box
        draw.ellipse((left + 25, top + 25, left + 105, top + 105), fill="#315E85")
        _text_center(draw, (left + 25, top + 25, left + 105, top + 105), num, _font(38, bold=True), fill="white")
        _text_center(draw, (left + 30, top + 125, right - 30, top + 245), heading, box_font)
        _text_center(draw, (left + 30, top + 250, right - 30, top + 355), detail, small_font, fill="#485563")
    for start, end in [((430, 410), (540, 410)), ((910, 410), (1020, 410)), ((1390, 410), (1500, 410))]:
        _arrow(draw, start, end)
    draw.text((325, 700), "候选召回缩小检索范围，精排与数量预测共同确定标准化结果", font=body_font, fill="#56616D")
    _save(image, "method_simple.png")

    # Figure 3-2: retrieval funnel.
    image, draw = _new_canvas(1800, 1040)
    draw.text((80, 50), "完整 ICD 词表", font=title_font, fill="#1F2933")
    _rounded_box(draw, (100, 135, 1700, 300), "37645 个标准疾病名称", fill="#E9EEF3", font=box_font)
    _rounded_box(draw, (160, 400, 790, 610), "字符 TF-IDF\n保留字面、缩写和部位", fill="#EAF2F8", font=body_font)
    _rounded_box(draw, (1010, 400, 1640, 610), "BGE 稠密召回\n补充同义和词序变化", fill="#DFF2E1", font=body_font)
    _rounded_box(draw, (560, 740, 1240, 950), "RRF 合并名次\n最终保留 Top-400", fill="#FFF3CD", font=box_font)
    _arrow(draw, (620, 300), (470, 400))
    _arrow(draw, (1180, 300), (1330, 400))
    _arrow(draw, (470, 610), (720, 740))
    _arrow(draw, (1330, 610), (1080, 740))
    _save(image, "retrieval_funnel.png")

    # Figure 3-3: exactly how the knowledge graph is used.
    image, draw = _new_canvas(1900, 1050)
    _rounded_box(draw, (650, 60, 1250, 240), "CPubMed-KG", fill="#EDE2F8", font=_font(46, bold=True))
    uses = [
        ((70, 410, 570, 735), "可靠一跳别名", "组成候选画像\n最多保留 3 个别名", "#EAF2F8"),
        ((700, 410, 1200, 735), "合成训练写法", "563 个额外训练组\n扩大可靠表达范围", "#DFF2E1"),
        ((1330, 410, 1830, 735), "关系难负例", "鉴别诊断、并发症\n用于区分相关实体", "#FDE2E4"),
    ]
    for box, heading, detail, fill in uses:
        draw.rounded_rectangle(box, radius=30, fill=fill, outline="#596777", width=4)
        left, top, right, bottom = box
        _text_center(draw, (left + 25, top + 30, right - 25, top + 145), heading, box_font)
        _text_center(draw, (left + 30, top + 155, right - 30, bottom - 25), detail, small_font, fill="#485563")
    for start, end in [((780, 240), (320, 410)), ((950, 240), (950, 410)), ((1120, 240), (1580, 410))]:
        _arrow(draw, start, end)
    draw.rounded_rectangle((340, 865, 1560, 1000), radius=25, fill="#F6F7F8", outline="#8994A0", width=3)
    _text_center(draw, (340, 865, 1560, 1000), "一跳别名用于候选表示，医学关联关系用于难负例构造", body_font, fill="#485563")
    _save(image, "kg_role.png")

    # Figure 3-4: count and Qwen only select the threshold group.
    image, draw = _new_canvas(1900, 1020)
    _rounded_box(draw, (610, 40, 1290, 210), "原始诊断短语", fill="#E9EEF3", font=box_font)
    _rounded_box(draw, (120, 360, 780, 590), "数量分类器", fill="#EAF2F8", font=box_font)
    _text_center(draw, (120, 485, 780, 570), "预测：1 个 / 2 个 / 3 个以上", small_font, fill="#485563")
    _rounded_box(draw, (1120, 360, 1780, 590), "Qwen 结构拆分", fill="#F3E8FF", font=box_font)
    _text_center(draw, (1120, 485, 1780, 570), "提取：单一成分 / 多个成分", small_font, fill="#485563")
    _rounded_box(draw, (600, 720, 1300, 890), "组合成 6 个解码组\n选择对应固定阈值", fill="#FFF3CD", font=box_font)
    _arrow(draw, (760, 210), (450, 360))
    _arrow(draw, (1140, 210), (1450, 360))
    _arrow(draw, (450, 590), (760, 720))
    _arrow(draw, (1450, 590), (1140, 720))
    draw.text((440, 930), "标准化结果由候选得分和固定词表共同确定", font=_font(34, bold=True), fill="#485563")
    _save(image, "count_qwen_decode.png")

    # Figure 3-5: training and frozen official evaluation.
    image, draw = _new_canvas(1900, 900)
    draw.text((90, 60), "训练阶段", font=title_font, fill="#1F5F8B")
    train = [
        ((70, 165, 420, 350), "6000 条\n官方训练数据", "#EAF2F8"),
        ((520, 165, 870, 350), "构造候选\n和图谱证据", "#DFF2E1"),
        ((970, 165, 1320, 350), "训练并保存\n13 个模型", "#FFF3CD"),
        ((1420, 165, 1770, 350), "固定版本、\n权重和阈值", "#F3E8FF"),
    ]
    for box, text, fill in train:
        _rounded_box(draw, box, text, fill=fill, font=body_font)
    for start, end in [((420, 255), (520, 255)), ((870, 255), (970, 255)), ((1320, 255), (1420, 255))]:
        _arrow(draw, start, end)
    draw.line((80, 485, 1820, 485), fill="#AAB4BF", width=4)
    draw.text((90, 535), "验证集评估阶段", font=title_font, fill="#2D6C45")
    evaluate = [
        ((160, 650, 590, 850), "2000 条\n官方验证数据", "#E8F5E9"),
        ((735, 650, 1165, 850), "使用固定参数\n完成模型推理", "#FFF3CD"),
        ((1310, 650, 1740, 850), "保存逐条预测\n独立复算 PASS", "#EAF2F8"),
    ]
    for box, text, fill in evaluate:
        _rounded_box(draw, box, text, fill=fill, font=body_font)
    _arrow(draw, (590, 750), (735, 750))
    _arrow(draw, (1165, 750), (1310, 750))
    _save(image, "training_inference.png")

    # Figure 4-1: retrieval recall curves.
    image, draw = _new_canvas(1800, 1040)
    left, top, right, bottom = 170, 85, 1700, 850
    ks = [50, 100, 200, 400, 800]
    series = [
        ("字符 TF-IDF", [69.68, 76.40, 81.35, 85.92, 88.61], "#4C78A8"),
        ("BGE", [84.07, 88.81, 91.27, 92.38, 93.26], "#E15759"),
        ("RRF 融合", [85.69, 89.46, 91.75, 92.66, 93.26], "#59A14F"),
    ]
    y_min, y_max = 65.0, 96.0
    x_pos = lambda i: left + i * (right - left) / (len(ks) - 1)
    y_pos = lambda value: bottom - (value - y_min) * (bottom - top) / (y_max - y_min)
    for tick in [65, 70, 75, 80, 85, 90, 95]:
        y = y_pos(tick)
        draw.line((left, y, right, y), fill="#D9DEE5", width=2)
        draw.text((95, y - 18), str(tick), font=small_font, fill="#3C4651")
    draw.line((left, top, left, bottom, right, bottom), fill="#46515C", width=4)
    for i, k in enumerate(ks):
        x = x_pos(i)
        draw.text((x - 25, bottom + 25), str(k), font=small_font, fill="#3C4651")
    for name, values, color in series:
        points = [(x_pos(i), y_pos(value)) for i, value in enumerate(values)]
        draw.line(points, fill=color, width=8, joint="curve")
        for x, y in points:
            draw.ellipse((x - 10, y - 10, x + 10, y + 10), fill="white", outline=color, width=6)
    for idx, (name, _, color) in enumerate(series):
        x = 700 + idx * 330
        draw.line((x, 35, x + 65, 35), fill=color, width=8)
        draw.text((x + 78, 14), name, font=small_font, fill="#26313C")
    draw.text((785, 930), "候选数量 K", font=body_font, fill="#26313C")
    y_label = Image.new("RGBA", (330, 90), (255, 255, 255, 0))
    y_draw = ImageDraw.Draw(y_label)
    y_draw.text((0, 0), "标签召回率（%）", font=body_font, fill="#26313C")
    y_label = y_label.rotate(90, expand=True)
    image.paste(y_label, (20, 350), y_label)
    _save(image, "retrieval_recall_v2.png")

    # Figure 4-2: performance across method stages.
    image, draw = _new_canvas(1800, 1030)
    labels = ["字符 TF-IDF\nTop-1", "基础文本\n匹配模型", "基础图谱\n增强方法", "本文完整\n方法"]
    values = [28.07, 52.92, 55.29, 71.73]
    colors = ["#9AA5B1", "#6FA8DC", "#8E7CC3", "#4F9D69"]
    left, top, right, bottom = 150, 90, 1700, 820
    y_max = 80
    for tick in range(0, 81, 10):
        y = bottom - tick / y_max * (bottom - top)
        draw.line((left, y, right, y), fill="#DEE3E8", width=2)
        draw.text((80, y - 18), str(tick), font=small_font, fill="#45515D")
    draw.line((left, top, left, bottom, right, bottom), fill="#45515D", width=4)
    slot = (right - left) / len(values)
    for i, (label, value, color) in enumerate(zip(labels, values, colors, strict=True)):
        center = left + slot * (i + 0.5)
        bar_w = slot * 0.55
        y = bottom - value / y_max * (bottom - top)
        draw.rounded_rectangle((center - bar_w / 2, y, center + bar_w / 2, bottom), radius=14, fill=color)
        _text_center(draw, (int(center - 100), int(y - 65), int(center + 100), int(y - 5)), f"{value:.2f}%", small_font)
        _text_center(draw, (int(center - 150), bottom + 20, int(center + 150), bottom + 130), label, small_font)
    draw.text((25, 20), "Micro-F1（%）", font=body_font, fill="#26313C")
    _save(image, "result_evolution.png")

    # Figure 4-3: fixed-parameter ablation.
    image, draw = _new_canvas(1900, 1140)
    labels = ["完整方法", "去掉 BGE 重排", "去掉图谱画像", "去掉 Qwen 分组", "去掉训练原型", "去掉普通 MacBERT", "去掉召回得分"]
    values = [71.73, 71.62, 70.43, 69.49, 69.25, 68.75, 67.02]
    left, top, right, bottom = 500, 80, 1800, 1050
    x_min, x_max = 65.0, 72.5
    x_pos = lambda value: left + (value - x_min) * (right - left) / (x_max - x_min)
    for tick in [65, 66, 67, 68, 69, 70, 71, 72]:
        x = x_pos(tick)
        draw.line((x, top, x, bottom), fill="#E0E4E8", width=2)
        draw.text((x - 16, bottom + 18), str(tick), font=small_font, fill="#45515D")
    row_h = (bottom - top) / len(labels)
    for i, (label, value) in enumerate(zip(labels, values, strict=True)):
        y = top + row_h * (i + 0.5)
        draw.text((55, y - 20), label, font=small_font, fill="#26313C")
        color = "#4F9D69" if i == 0 else "#6FA8DC"
        draw.rounded_rectangle((left, y - 26, x_pos(value), y + 26), radius=12, fill=color)
        draw.text((x_pos(value) + 18, y - 20), f"{value:.2f}", font=small_font, fill="#26313C")
    draw.text((820, 1085), "Micro-F1（%）", font=body_font, fill="#26313C")
    _save(image, "ablation_effects.png")

    # Figure 4-4: subgroup performance.
    image, draw = _new_canvas(1800, 1040)
    labels = ["总体", "ICD 内", "多实体", "单实体", "结构激活", "未激活"]
    values = [71.73, 75.08, 78.77, 55.63, 73.23, 69.94]
    colors = ["#4C78A8", "#72B7B2", "#59A14F", "#E15759", "#B279A2", "#9D755D"]
    left, top, right, bottom = 150, 80, 1700, 820
    y_min, y_max = 45, 82
    y_pos = lambda value: bottom - (value - y_min) * (bottom - top) / (y_max - y_min)
    for tick in [45, 50, 55, 60, 65, 70, 75, 80]:
        y = y_pos(tick)
        draw.line((left, y, right, y), fill="#DEE3E8", width=2)
        draw.text((80, y - 18), str(tick), font=small_font, fill="#45515D")
    draw.line((left, top, left, bottom, right, bottom), fill="#45515D", width=4)
    slot = (right - left) / len(values)
    for i, (label, value, color) in enumerate(zip(labels, values, colors, strict=True)):
        center = left + slot * (i + 0.5)
        bar_w = slot * 0.58
        y = y_pos(value)
        draw.rounded_rectangle((center - bar_w / 2, y, center + bar_w / 2, bottom), radius=12, fill=color)
        draw.text((center - 48, y - 45), f"{value:.2f}", font=small_font, fill="#26313C")
        _text_center(draw, (int(center - 110), bottom + 18, int(center + 110), bottom + 95), label, small_font)
    draw.text((25, 20), "Micro-F1（%）", font=body_font, fill="#26313C")
    _save(image, "subgroup_performance_v2.png")

    # Figure 5-1: relation between normalization and symptom-based inference.
    image, draw = _new_canvas(1900, 810)
    _rounded_box(draw, (80, 180, 820, 700), "疾病名称标准化\n\n输入：病案中的疾病名称\n例如：肺部恶性肿瘤\n\n输出：统一 ICD 名称和编码", fill="#EAF2F8", font=body_font)
    _rounded_box(draw, (1080, 180, 1820, 700), "症状辅助诊断\n\n输入：症状、病史和检查信息\n例如：咳嗽、夜间加重\n\n输出：候选疾病", fill="#FDEBD0", font=body_font)
    _arrow(draw, (820, 440), (1080, 440), fill="#7C6A45", width=8)
    draw.text((830, 330), "名称与编码标准化", font=small_font, fill="#7C5428")
    draw.text((330, 735), "症状推断给出候选疾病，标准化模块统一疾病名称和编码", font=title_font, fill="#38434F")
    _save(image, "future_roadmap.png")


def add_toc_v2(doc) -> None:
    title = doc.add_paragraph()
    title.alignment = base.WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    run = title.add_run("目　录")
    base.set_run_font(run, east_asia="Heiti SC", size=16, bold=True)

    entries = [
        ("摘要", "i", 0),
        ("Abstract", "ii", 0),
        ("第1章 绪论", "1", 0),
        ("1.1 研究背景与意义", "1", 1),
        ("1.2 研究问题定义", "2", 1),
        ("1.3 研究内容与论文结构", "3", 1),
        ("第2章 相关研究与数据", "4", 0),
        ("2.1 疾病实体标准化研究进展", "4", 1),
        ("2.2 数据集与知识资源", "5", 1),
        ("2.3 数据特征与任务难点", "6", 1),
        ("第3章 疾病实体标准化方法", "8", 0),
        ("3.1 方法总体框架", "8", 1),
        ("3.2 多路候选实体召回", "9", 1),
        ("3.3 多源特征融合与候选精排", "10", 1),
        ("3.4 实体数量预测与约束解码", "12", 1),
        ("3.5 模型训练与复现实验", "13", 1),
        ("第4章 实验结果与分析", "15", 0),
        ("4.1 实验设置与评价指标", "15", 1),
        ("4.2 候选召回性能分析", "15", 1),
        ("4.3 整体性能分析", "16", 1),
        ("4.4 消融实验与模块贡献分析", "18", 1),
        ("4.5 分组性能与误差分析", "19", 1),
        ("第5章 结论与展望", "21", 0),
        ("5.1 主要研究结论", "21", 1),
        ("5.2 研究局限与展望", "21", 1),
        ("参考文献", "23", 0),
        ("致谢", "27", 0),
        ("附录 主要实验参数与复现配置", "28", 0),
    ]
    for entry_title, page, level in entries:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.74 if level else 0)
        p.paragraph_format.space_before = Pt(3 if not level else 0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p_pr = p._p.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:leader"), "dot")
        tab.set(qn("w:pos"), str(base.TOKENS["content_width_dxa"] - (420 if level else 0)))
        tabs.append(tab)
        p_pr.append(tabs)
        title_run = p.add_run(entry_title)
        base.set_run_font(title_run, size=10.5, bold=not level)
        tab_run = p.add_run("\t")
        base.set_run_font(tab_run, size=10.5)
        page_run = p.add_run(page)
        base.set_run_font(page_run, size=10.5, bold=not level)


base.generate_figures = generate_figures_v2
base.add_toc = add_toc_v2


if __name__ == "__main__":
    print(base.build())
